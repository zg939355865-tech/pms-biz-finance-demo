from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


BASE_DIR = Path(__file__).resolve().parent
SOURCE_DOCX = BASE_DIR.parent / "docs" / "PRD" / "UR（昕彤赋能PMS业财一体化）项目管理需求规格说明书-2026年-03月-02日-v0.1.docx"
OUTPUT_DOCX = BASE_DIR / "PMS项目管理需求规格说明书_标准模板_V1.1.docx"

NAVY = "16324F"
BLUE = "246BCE"
TEAL = "18A39B"
TEXT = "25364A"
MUTED = "6B7A90"
LIGHT = "F5F7FA"
LIGHT_BLUE = "EDF4F8"
BORDER = "D7E0EA"
WHITE = "FFFFFF"
FONT_CN = "微软雅黑"
FONT_LATIN = "Arial"


def set_run_font(run, name=FONT_CN, size=10.5, color=TEXT, bold=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.get_or_add_rFonts()
    r_fonts.set(qn("w:ascii"), FONT_LATIN)
    r_fonts.set(qn("w:hAnsi"), FONT_LATIN)
    r_fonts.set(qn("w:eastAsia"), name)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=90, start=130, bottom=90, end=130):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        tag = tc_mar.find(qn(f"w:{edge}"))
        if tag is None:
            tag = OxmlElement(f"w:{edge}")
            tc_mar.append(tag)
        tag.set(qn("w:w"), str(value))
        tag.set(qn("w:type"), "dxa")


def set_table_borders(table, color=BORDER, size="6"):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:color"), color)


def remove_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")


def set_paragraph_border(paragraph, edge="bottom", color=BLUE, size="10", space="5"):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.find(qn("w:pBdr"))
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    border = p_bdr.find(qn(f"w:{edge}"))
    if border is None:
        border = OxmlElement(f"w:{edge}")
        p_bdr.append(border)
    border.set(qn("w:val"), "single")
    border.set(qn("w:sz"), size)
    border.set(qn("w:space"), space)
    border.set(qn("w:color"), color)


def set_paragraph_shading(paragraph, fill):
    p_pr = paragraph._p.get_or_add_pPr()
    shading = p_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        p_pr.append(shading)
    shading.set(qn("w:val"), "clear")
    shading.set(qn("w:fill"), fill)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = tr_pr.find(qn("w:tblHeader"))
    if tbl_header is None:
        tbl_header = OxmlElement("w:tblHeader")
        tr_pr.append(tbl_header)
    tbl_header.set(qn("w:val"), "true")


def remove_page_border(section):
    sect_pr = section._sectPr
    pg_borders = sect_pr.find(qn("w:pgBorders"))
    if pg_borders is not None:
        sect_pr.remove(pg_borders)


def clear_story(story):
    element = story._element
    for child in list(element):
        element.remove(child)
    story.add_paragraph()


def add_page_field(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_sep = OxmlElement("w:fldChar")
    fld_sep.set(qn("w:fldCharType"), "separate")
    value = OxmlElement("w:t")
    value.text = "1"
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    for node in (fld_begin, instr, fld_sep, value, fld_end):
        run._r.append(node)
    set_run_font(run, size=8.5, color=MUTED)


def configure_sections(doc):
    for index, section in enumerate(doc.sections):
        section.page_width = Cm(21)
        section.page_height = Cm(29.7)
        section.top_margin = Cm(1.75 if index else 1.35)
        section.bottom_margin = Cm(1.65 if index else 1.35)
        section.left_margin = Cm(2.05)
        section.right_margin = Cm(2.05)
        section.header_distance = Cm(0.75)
        section.footer_distance = Cm(0.7)
        remove_page_border(section)

    cover = doc.sections[0]
    cover.different_first_page_header_footer = True
    cover.first_page_header.is_linked_to_previous = False
    cover.first_page_footer.is_linked_to_previous = False
    clear_story(cover.first_page_header)
    clear_story(cover.first_page_footer)

    for section in doc.sections[1:]:
        section.header.is_linked_to_previous = False
        section.footer.is_linked_to_previous = False
        clear_story(section.header)
        clear_story(section.footer)
        usable_width = section.page_width - section.left_margin - section.right_margin

        header_table = section.header.add_table(rows=1, cols=2, width=usable_width)
        header_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        header_table.autofit = False
        remove_table_borders(header_table)
        left, right = header_table.rows[0].cells
        left.width = Cm(11.6)
        right.width = Cm(5.3)
        left.text = "PMS 业财一体化  /  PRODUCT REQUIREMENTS"
        right.text = "需求规格说明书  ·  L0"
        left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
        for cell in (left, right):
            set_cell_margins(cell, 30, 0, 90, 0)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=8, color=NAVY, bold=cell is left)
        set_paragraph_border(left.paragraphs[0], color="BFD0E0", size="5", space="4")
        set_paragraph_border(right.paragraphs[0], color="BFD0E0", size="5", space="4")

        footer_table = section.footer.add_table(rows=1, cols=2, width=usable_width)
        footer_table.alignment = WD_TABLE_ALIGNMENT.CENTER
        footer_table.autofit = False
        remove_table_borders(footer_table)
        f_left, f_right = footer_table.rows[0].cells
        f_left.text = "昕彤赋能  |  PMS产品与交付中心"
        f_left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in f_left.paragraphs[0].runs:
            set_run_font(run, size=8.5, color=MUTED)
        add_page_field(f_right.paragraphs[0])


def configure_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = FONT_CN
    normal.font.size = Pt(10.2)
    normal.font.color.rgb = RGBColor.from_string(TEXT)
    normal._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
    pf = normal.paragraph_format
    pf.line_spacing = 1.28
    pf.space_after = Pt(4)
    pf.widow_control = True

    style_roles = {
        "一级标题": (17, NAVY, 18, 9),
        "Heading 1": (17, NAVY, 18, 9),
        "二级标题": (13.5, BLUE, 13, 6),
        "Heading 2": (13.5, BLUE, 13, 6),
        "三级标题": (11, NAVY, 9, 4),
        "Heading 3": (11, NAVY, 9, 4),
    }
    for name, (size, color, before, after) in style_roles.items():
        if name not in doc.styles:
            continue
        style = doc.styles[name]
        style.font.name = FONT_CN
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:hAnsi"), FONT_LATIN)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), FONT_CN)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True


def style_existing_paragraphs(doc):
    h1 = {"一级标题", "Heading 1"}
    h2 = {"二级标题", "Heading 2"}
    h3 = {"三级标题", "Heading 3"}
    for paragraph in doc.paragraphs:
        name = paragraph.style.name
        text = paragraph.text.strip()
        has_image = bool(paragraph._p.xpath(".//w:drawing | .//w:pict"))
        fmt = paragraph.paragraph_format

        if text == "本文档记录项目管理相关功能。":
            paragraph.style = doc.styles["Normal"]
            fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            fmt.first_line_indent = Cm(0.74)
            fmt.line_spacing = 1.28
            fmt.space_after = Pt(4)
            for run in paragraph.runs:
                set_run_font(run, size=10.2, color=TEXT)
        elif name in h1:
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            set_paragraph_border(paragraph, edge="left", color=TEAL, size="20", space="8")
            set_paragraph_shading(paragraph, "F7FAFC")
            for run in paragraph.runs:
                set_run_font(run, size=17, color=NAVY, bold=True)
        elif name in h2:
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                set_run_font(run, size=13.5, color=BLUE, bold=True)
        elif name in h3:
            fmt.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in paragraph.runs:
                set_run_font(run, size=11, color=NAVY, bold=True)
        elif name.startswith("TOC") or name.startswith("目录"):
            fmt.space_after = Pt(3)
            for run in paragraph.runs:
                set_run_font(run, size=9.5, color=TEXT)
        elif text.startswith("图") or text.startswith("表") and len(text) < 36:
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.space_before = Pt(4)
            fmt.space_after = Pt(8)
            for run in paragraph.runs:
                set_run_font(run, size=9, color=MUTED)
        elif has_image:
            fmt.alignment = WD_ALIGN_PARAGRAPH.CENTER
            fmt.space_before = Pt(6)
            fmt.space_after = Pt(6)
        else:
            if text:
                fmt.line_spacing = 1.28
                fmt.space_after = Pt(4)
                if not paragraph._p.xpath(".//w:numPr"):
                    fmt.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            for run in paragraph.runs:
                set_run_font(run, size=10.2, color=TEXT)


def style_existing_tables(doc):
    for table_index, table in enumerate(doc.tables):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        set_table_borders(table, color=BORDER, size="4")
        if table.rows:
            set_repeat_table_header(table.rows[0])

        for row_index, row in enumerate(table.rows):
            for cell_index, cell in enumerate(row.cells):
                cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                set_cell_margins(cell)
                if row_index == 0:
                    set_cell_shading(cell, NAVY)
                elif cell_index == 0 and len(row.cells) <= 6:
                    set_cell_shading(cell, LIGHT_BLUE)
                elif row_index % 2 == 0:
                    set_cell_shading(cell, "FAFCFD")

                for paragraph in cell.paragraphs:
                    paragraph.paragraph_format.space_before = Pt(0)
                    paragraph.paragraph_format.space_after = Pt(0)
                    paragraph.paragraph_format.line_spacing = 1.15
                    if row_index == 0:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for run in paragraph.runs:
                        set_run_font(
                            run,
                            size=8.8,
                            color=WHITE if row_index == 0 else TEXT,
                            bold=True if row_index == 0 or cell_index == 0 else run.bold,
                        )


def remove_old_cover(doc):
    paragraphs = doc.paragraphs
    anchor_index = next((index for index, paragraph in enumerate(paragraphs) if paragraph.text.strip() == "文件信息"), None)
    if anchor_index is None:
        raise RuntimeError("Cannot locate the first content page anchor.")
    # Keep the final blank paragraph that owns the cover section break.
    for paragraph in list(paragraphs[: max(0, anchor_index - 1)]):
        paragraph._element.getparent().remove(paragraph._element)
    return doc.paragraphs[0]


def insert_table_before(doc, anchor, rows, cols):
    table = doc.add_table(rows=rows, cols=cols)
    anchor._p.addprevious(table._tbl)
    return table


def add_cover(doc, anchor):
    brand = insert_table_before(doc, anchor, 1, 2)
    brand.alignment = WD_TABLE_ALIGNMENT.CENTER
    brand.autofit = False
    remove_table_borders(brand)
    brand_left, brand_right = brand.rows[0].cells
    brand_left.width = Cm(10.9)
    brand_right.width = Cm(5.9)
    for cell in (brand_left, brand_right):
        set_cell_shading(cell, NAVY)
        set_cell_margins(cell, 150, 210, 150, 210)
    brand_left.text = "DAWN  昕彤赋能"
    brand_right.text = "PMS  /  CONTROLLED DOCUMENT"
    brand_left.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
    brand_right.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in brand_left.paragraphs[0].runs:
        set_run_font(run, size=11.5, color=WHITE, bold=True)
    for run in brand_right.paragraphs[0].runs:
        set_run_font(run, size=8, color="DCE8F5", bold=True)

    spacer = anchor.insert_paragraph_before("")
    spacer.paragraph_format.space_after = Pt(58)

    eyebrow = anchor.insert_paragraph_before("PMS 业财一体化系统  ·  标准需求文档")
    eyebrow.alignment = WD_ALIGN_PARAGRAPH.LEFT
    eyebrow.paragraph_format.space_after = Pt(12)
    for run in eyebrow.runs:
        set_run_font(run, size=10.5, color=TEAL, bold=True)

    title = anchor.insert_paragraph_before("项目管理需求规格说明书")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    title.paragraph_format.space_after = Pt(8)
    for run in title.runs:
        set_run_font(run, size=29, color=NAVY, bold=True)

    subtitle = anchor.insert_paragraph_before("PRODUCT REQUIREMENTS DOCUMENT  /  PRD · UR")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    subtitle.paragraph_format.space_after = Pt(8)
    for run in subtitle.runs:
        set_run_font(run, size=9.2, color=MUTED, bold=True)

    title_rule = anchor.insert_paragraph_before("")
    title_rule.paragraph_format.space_after = Pt(18)
    set_paragraph_border(title_rule, edge="bottom", color=TEAL, size="18", space="2")

    lead = anchor.insert_paragraph_before("面向PMS项目需求分析、产品设计、实施交付与验收的统一需求基线。")
    lead.alignment = WD_ALIGN_PARAGRAPH.LEFT
    lead.paragraph_format.space_after = Pt(24)
    for run in lead.runs:
        set_run_font(run, size=10.2, color=MUTED)

    meta = insert_table_before(doc, anchor, 2, 4)
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta.autofit = False
    set_table_borders(meta, color=BORDER, size="4")
    meta_values = [
        ("文档类型", "PRD / UR", "适用系统", "PMS业财一体化系统"),
        ("版本状态", "V1.1 · 标准模板", "发布日期", "YYYY-MM-DD"),
    ]
    for row, values in zip(meta.rows, meta_values):
        widths = (Cm(2.5), Cm(5.7), Cm(2.5), Cm(6.1))
        for index, (cell, width, value) in enumerate(zip(row.cells, widths, values)):
            cell.width = width
            cell.text = value
            set_cell_margins(cell, 125, 150, 125, 150)
            is_label = index in (0, 2)
            set_cell_shading(cell, LIGHT_BLUE if is_label else WHITE)
            for run in cell.paragraphs[0].runs:
                set_run_font(run, size=8.8, color=NAVY if is_label else TEXT, bold=is_label)

    approval_label = anchor.insert_paragraph_before("DOCUMENT CONTROL")
    approval_label.paragraph_format.space_before = Pt(34)
    approval_label.paragraph_format.space_after = Pt(8)
    for run in approval_label.runs:
        set_run_font(run, size=8.5, color=TEAL, bold=True)

    approval = insert_table_before(doc, anchor, 1, 3)
    approval.alignment = WD_TABLE_ALIGNMENT.CENTER
    approval.autofit = False
    set_table_borders(approval, color=BORDER, size="4")
    for index, label in enumerate(("编写", "审核", "批准")):
        cell = approval.rows[0].cells[index]
        cell.text = f"{label}\n姓名：________\n日期：____-__-__"
        set_cell_shading(cell, "F8FAFC")
        set_cell_margins(cell, 140, 160, 140, 160)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run_index, run in enumerate(cell.paragraphs[0].runs):
            set_run_font(run, size=8.5, color=NAVY if run_index == 0 else MUTED, bold=run_index == 0)

    footer_note = anchor.insert_paragraph_before("内部受控文档  ·  未经授权请勿外传")
    footer_note.alignment = WD_ALIGN_PARAGRAPH.LEFT
    footer_note.paragraph_format.space_before = Pt(36)
    for run in footer_note.runs:
        set_run_font(run, size=8.2, color=MUTED)


def set_update_fields(doc):
    settings = doc.settings._element
    update_fields = settings.find(qn("w:updateFields"))
    if update_fields is None:
        update_fields = OxmlElement("w:updateFields")
        settings.append(update_fields)
    # Keep opening fast and predictable. The build/QA flow refreshes fields
    # explicitly in Word before publishing the final template.
    update_fields.set(qn("w:val"), "false")


def main():
    if not SOURCE_DOCX.exists():
        raise FileNotFoundError(SOURCE_DOCX)

    doc = Document(SOURCE_DOCX)
    doc.core_properties.title = "PMS项目管理需求规格说明书标准模板"
    doc.core_properties.subject = "PMS PRD/UR标准模板"
    doc.core_properties.comments = "现代企业规范版；由原项目管理需求规格说明书优化。"

    configure_styles(doc)
    style_existing_paragraphs(doc)
    style_existing_tables(doc)
    anchor = remove_old_cover(doc)
    add_cover(doc, anchor)
    configure_sections(doc)
    set_update_fields(doc)

    OUTPUT_DOCX.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT_DOCX)
    print(OUTPUT_DOCX)


if __name__ == "__main__":
    main()
